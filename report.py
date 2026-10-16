"""Structured report object and markdown renderer."""
from __future__ import annotations

from dataclasses import dataclass, field


@dataclass
class Report:
    technique: str = ""
    findings: list[str] = field(default_factory=list)
    measurements: dict = field(default_factory=dict)
    impression: str = ""

    @classmethod
    def from_dict(cls, d: dict) -> "Report":
        return cls(
            technique=d.get("technique", ""),
            findings=list(d.get("findings", [])),
            measurements=dict(d.get("measurements", {})),
            impression=d.get("impression", ""),
        )

    def to_markdown(self) -> str:
        # No document title: the panel already says "Befundentwurf" — repeating it inside
        # the text was redundant. "Befund" renders only when the model produced findings;
        # the morphometry module deliberately leaves it empty.
        lines = [f"**Technik:** {self.technique or '—'}", ""]

        if self.findings:
            lines += ["## Befund"]
            lines += [f"- {f}" for f in self.findings]
            lines += [""]

        if self.measurements:
            lines += ["## Morphometrie"]
            lines += [f"- **{k}:** {v}" for k, v in self.measurements.items()]
            lines += [""]

        lines += ["## Beurteilung", self.impression or "—"]
        return "\n".join(lines)

    def to_docx_bytes(self) -> bytes:
        """Render the report as an editable Word (.docx) document."""
        import io

        from docx import Document

        doc = Document()
        doc.add_heading("Befundentwurf – Morphometrie", level=0)

        p = doc.add_paragraph()
        p.add_run("Technik: ").bold = True
        p.add_run(self.technique or "—")

        if self.findings:
            doc.add_heading("Befund", level=1)
            for f in self.findings:
                doc.add_paragraph(f, style="List Bullet")

        if self.measurements:
            doc.add_heading("Morphometrie", level=1)
            for k, v in self.measurements.items():
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{k}: ").bold = True
                p.add_run(str(v))

        doc.add_heading("Beurteilung", level=1)
        doc.add_paragraph(self.impression or "—")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
